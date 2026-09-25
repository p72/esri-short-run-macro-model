"""ESRI 2015年基準 SNA ファイルから四半期系列を作る.

入力（data/raw/）:
  gaku-jk2522.csv / gaku-mk2522.csv / def-qk2522.csv  四半期GDP速報（2025年4-6月期2次、季調・年率）
  2022qom2_jp.xlsx   国民所得・国民可処分所得の分配（四半期・原系列）
  2022i4_jp.xlsx     一般政府 所得支出勘定（四半期シート）
  2022i5_jp.xlsx     家計 所得支出勘定（四半期シート）
出力: data/processed/sna_quarterly.csv（モデル変数名、10億円・年率、デフレーターは2015年=1）
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import vintage as VT  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
OUT = VT.processed("sna_quarterly.csv")


# ---------------------------------------------------------------------------
# 四半期GDP速報 CSV
# ---------------------------------------------------------------------------
def read_qe(fname: str, cols: dict[str, int]) -> pd.DataFrame:
    raw = VT.qe_file(fname).read_bytes().decode("shift_jis", errors="replace")
    df = pd.read_csv(io.StringIO(raw), header=None)
    periods, rows = [], []
    year = None
    qmap = {"1- 3": 1, "4- 6": 2, "7- 9": 3, "10-12": 4}
    for _, r in df.iloc[7:].iterrows():
        label = str(r[0]).strip()
        if label[:4].isdigit() and label[4:5] == "/":
            year = int(label[:4])
        q = next((v for k, v in qmap.items() if k in label), None)
        if q is None or year is None:
            continue
        periods.append(pd.Period(year=year, quarter=q, freq="Q"))
        rows.append({name: pd.to_numeric(str(r[c]).replace(",", "").strip(), errors="coerce")
                     for name, c in cols.items()})
    return pd.DataFrame(rows, index=pd.PeriodIndex(periods, freq="Q"))


def official_ywv() -> pd.Series | None:
    """QE 表 kshotoku（雇用者報酬）の名目季節調整系列（年率、10億円）。無い版では None."""
    f = VT.qe_file("kshotoku-q")
    if not f.exists():
        return None
    raw = f.read_bytes().decode("shift_jis", errors="replace")
    df = pd.read_csv(io.StringIO(raw), header=None)
    periods, vals, year = [], [], None
    qmap = {"1- 3": 1, "4- 6": 2, "7- 9": 3, "10-12": 4}
    for _, r in df.iloc[8:].iterrows():
        label = str(r[0]).strip()
        if label[:4].isdigit() and label[4:5] == "/":
            year = int(label[:4])
        q = next((v for k, v in qmap.items() if k in label), None)
        if q is None or year is None:
            continue
        periods.append(pd.Period(year=year, quarter=q, freq="Q"))
        vals.append(pd.to_numeric(str(r[3]).replace(",", "").strip(), errors="coerce"))
    return pd.Series(vals, index=pd.PeriodIndex(periods, freq="Q"))


def smooth_annual(x: pd.Series) -> pd.Series:
    """暦年平均を保ちつつ四半期を滑らかにする（年平均を年央に置き3次スプラインで補間、年平均で再調整）."""
    from scipy.interpolate import CubicSpline
    x = x.dropna()
    years = sorted({p.year for p in x.index if (x.index.year == p.year).sum() == 4})
    ymean = pd.Series({y: x[[p.year == y for p in x.index]].mean() for y in years})
    t_mid = np.array([y + 0.5 for y in years])
    cs = CubicSpline(t_mid, ymean.values, bc_type="natural")
    t = np.array([p.year + (p.quarter - 0.5) / 4 for p in x.index])
    out = pd.Series(cs(t), index=x.index)
    for y in years:  # 補間後の年平均を元の年平均に一致させる
        m = [p.year == y for p in x.index]
        out[m] = out[m] * ymean[y] / out[m].mean()
    return out


def qe_block() -> pd.DataFrame:
    real = read_qe("gaku-jk", {
        "GDP": 1, "CP": 2, "IHP": 5, "IFP": 6, "INP": 7, "CG": 8, "IG": 9, "ING": 10,
        "BF": 11, "XGS": 12, "MGS": 13, "KAISA": 14})
    nom = read_qe("gaku-mk", {
        "GDPV": 1, "CPV": 2, "IHPV": 5, "IFPV": 6, "INPV": 7, "CGV": 8, "IGV": 9, "INGV": 10,
        "BFV": 11, "XGSV": 12, "MGSV": 13, "NFIV": 15, "RTRIV": 16, "PTRIV": 17, "GNIV": 18})
    defl = read_qe("def-qk", {
        "PGDP": 1, "PCP": 2, "PIHP": 5, "PIFP": 6, "PCG": 8, "PIG": 9, "PXGS": 12, "PMGS": 13}) / 100
    return real.join(nom).join(defl)


# ---------------------------------------------------------------------------
# 年次推計 Excel の四半期シート（行番号は 0 始まり、列1以降が1994Q1〜）
# ---------------------------------------------------------------------------
def read_q_sheet(fname: str, sheet: str, rows: dict[str, tuple[int, str]]) -> pd.DataFrame:
    ws = openpyxl.load_workbook(VT.annual_file(fname), read_only=True, data_only=True)[sheet]
    table = list(ws.iter_rows(values_only=True))
    ncol = len(table[6]) - 1
    idx = pd.period_range("1994Q1", periods=ncol, freq="Q")
    out = {}
    for name, (i, must) in rows.items():
        label = str(table[i][0]).replace("　", "").strip()
        if must not in label:
            raise ValueError(f"{fname}[{sheet}] 行{i} は「{must}」のはずが「{label}」")
        vals = [pd.to_numeric(v, errors="coerce") if v not in ("-", None) else np.nan for v in table[i][1:]]
        out[name] = vals
    return pd.DataFrame(out, index=idx)


def income_block() -> pd.DataFrame:
    qom = read_q_sheet("qom2", "実数", {
        "YWV": (7, "雇用者報酬"), "YWIV": (8, "賃金・俸給"), "YOLIV": (9, "雇主の社会負担"),
        "YIGV": (15, "一般政府"), "YIEV": (26, "家計"), "YINPV": (33, "対家計民間非営利団体"),
        "YICV": (49, "個人企業"), "NIV": (53, "国民所得（要素費用表示）"),
        "TAXNETV": (55, "生産・輸入品に課される税"),
        "ENTV": (42, "企業所得"),  # 企業部門の第１次所得バランス（法人＋公的企業＋個人企業）
        "RENTHH": (32, "賃貸料"),  # 家計財産所得のうち賃貸料（受取）
    })
    g1 = read_q_sheet("i4", "四半期（１）", {
        "CCAVG": (12, "固定資本減耗"), "TPIV": (16, "生産・輸入品に課される税"),
        "TCIV": (18, "付加価値型税"), "TCSTV": (19, "輸入関税"), "SUBV": (22, "補助金"),
    })
    g2 = read_q_sheet("i4", "四半期（２）", {
        "TINCGV": (25, "所得・富等に課される経常税"), "CSSV": (28, "純社会負担"),
    })
    h2 = read_q_sheet("i5", "四半期（２）", {
        "TYPV": (7, "所得・富等に課される経常税"), "YDV": (21, "可処分所得（純）"),
        "BSSV": (30, "現物社会移転以外の社会給付"),
    })
    df = qom.join(g1).join(g2).join(h2)
    df["TYCV"] = df["TINCGV"] - df["TYPV"]
    df["OITAXV"] = df["TPIV"] - df["TCIV"] - df["TCSTV"]
    return df


# ---------------------------------------------------------------------------
# 季節調整（移動平均比率法）
# ---------------------------------------------------------------------------
def seasonal_adjust(s: pd.Series, additive: bool = False, years: int = 5) -> pd.Series:
    """2×4 中心化移動平均からの比率(差)を四半期別に移動平均して季節因子とする."""
    x = s.astype(float)
    trend = (0.5 * x.shift(2) + x.shift(1) + x + x.shift(-1) + 0.5 * x.shift(-2)) / 4
    irr = (x - trend) if additive else (x / trend)
    q = pd.Series([p.quarter for p in x.index], index=x.index)
    factor = pd.Series(np.nan, index=x.index)
    for k in range(1, 5):
        sub = irr[q == k]
        sm = sub.rolling(years, center=True, min_periods=3).mean().ffill().bfill()
        factor[q == k] = sm
    norm = factor.rolling(4, center=True, min_periods=2).mean()
    factor = (factor - norm) if additive else (factor / norm)
    return (x - factor) if additive else (x / factor)


ADDITIVE = {"YIGV", "YINPV", "SUBV"}  # 符号が変わる、またはゼロ近傍を取りうる系列


def build() -> pd.DataFrame:
    qe = qe_block()
    inc = income_block()
    sa = pd.DataFrame({c: seasonal_adjust(inc[c].dropna(), additive=c in ADDITIVE).reindex(inc.index) * 4
                       for c in inc.columns})
    # 企業所得は変動が大きく加法型が適する（論文乗数表から逆算した法人企業所得の経路との形の誤差 2.5% vs 乗法 3.9%）
    for c in ["ENTV", "YICV", "RENTHH", "YINPV"]:
        sa[c] = seasonal_adjust(inc[c].dropna(), additive=True).reindex(inc.index) * 4
    df = qe.join(sa, how="left")
    ywv = official_ywv()
    if ywv is not None:  # 雇用者報酬は QE の公式季節調整値を優先（自前の季調との差は 0.15% 程度）
        off = ywv.reindex(df.index)
        df["YWV"] = off.where(off.notna(), df["YWV"])
    # 国民所得の定義式(83): 固定資本減耗(+不突合)は季調残差が乗らないよう年平均を滑らかに補間したものを使い、
    # NIV は定義式から求める（NIV を独立に季調すると法人企業所得（残差）に ±3% の季節性が残る）
    df["ITAXV"] = df["TCIV"] + df["TCSTV"] + df["OITAXV"]
    ccav_raw = df["GDPV"] - df["ITAXV"] + df["SUBV"] + df["NFIV"] - df["NIV"]
    df["CCAV"] = smooth_annual(ccav_raw).reindex(df.index)
    df["TAXV"] = df["TYPV"] + df["TYCV"] + df["ITAXV"]
    # 家計財産所得 YIEV は賃貸料を含まない金融的な財産所得（利子・配当・その他の投資所得）。
    # 論文の乗数表から逆算した法人企業所得の水準は、SNA 法人企業所得＋家計賃貸料＋NPISH財産所得 と 0.2% で一致し、
    # 式104（財産所得は金利に反応）の定式化とも整合する。
    df["YIEV"] = df["YIEV"] - df["RENTHH"]
    df["YIV"] = df["YIEV"] + df["YIGV"]
    # 法人企業所得は SNA の企業所得（法人＋公的）＋家計賃貸料＋NPISH財産所得 を直接季調したもの。
    # ESRI の四半期経路（乗数表から逆算）はこの系列と相関 0.96（国民所得の残差として作ると 0.61）。
    # 定義式(83)(84)の残りは不突合 SDV で閉じる
    df["YCV"] = df["ENTV"] - df["YICV"] + df["RENTHH"] + df["YINPV"]
    df["NIV"] = df["YWV"] + df["YIV"] + df["YICV"] + df["YCV"]
    df["SDV"] = df["GDPV"] - df["CCAV"] - df["ITAXV"] + df["SUBV"] + df["NFIV"] - df["NIV"]
    df["OTYDV"] = df["YDV"] - (df["YWV"] + df["BSSV"] + df["YIEV"] + df["YICV"] - df["TYPV"] - df["CSSV"])
    return df


# ---------------------------------------------------------------------------
# 年末ストック（固定資本ストック、在庫、対外資産負債）と政府の固定資本形成
# ---------------------------------------------------------------------------
_Z2H = str.maketrans("０１２３４５６７８９", "0123456789")


def sheet_year(name: str) -> int | None:
    z = name.translate(_Z2H)
    if z.startswith("平成") and z[2:].isdigit():
        return 1988 + int(z[2:])
    if z == "令和元":
        return 2019
    if z.startswith("令和") and z[2:].isdigit():
        return 2018 + int(z[2:])
    return None


def _find(rows, text: str, start: int = 0) -> tuple:
    for r in rows[start:]:
        if r and r[0] is not None and text in str(r[0]).replace("　", ""):
            return r
    raise KeyError(text)


def annual_block() -> pd.DataFrame:
    rec: dict[int, dict[str, float]] = {}
    # 固定資本ストック（名目、暦年末）: 列 3=非金融法人民間 4=同公的 6=金融機関民間 7=同公的 8=一般政府 9=家計 10=NPISH
    wb = openpyxl.load_workbook(VT.annual_file("ss4n"), read_only=True, data_only=True)
    for sn in wb.sheetnames:
        y = sheet_year(sn)
        if y is None:
            continue
        rows = list(wb[sn].iter_rows(values_only=True))
        house, total = _find(rows, "住宅"), _find(rows, "固定資産合計")
        f = lambda r, cols: sum(float(r[c] or 0) for c in cols)  # noqa: E731
        priv = (3, 6, 9, 10)
        rec.setdefault(y, {})["KHPV"] = f(house, priv)
        rec[y]["KFPV"] = f(total, priv) - rec[y]["KHPV"]
        rec[y]["KGV"] = f(total, (4, 7, 8))
    # 国民資産・負債残高（暦年末、列8=当期末残高）
    wb = openpyxl.load_workbook(VT.annual_file("ss1"), read_only=True, data_only=True)
    for sn in wb.sheetnames:
        y = sheet_year(sn)
        if y is None:
            continue
        rows = list(wb[sn].iter_rows(values_only=True))
        rec.setdefault(y, {})["KNPV"] = float(_find(rows, "ｂ．在庫")[8])
        rec[y]["LANDT"] = float(_find(rows, "（ａ）土地")[8])
        # 株式総額は負債側（居住者が発行した株式の残高）。資産側（居住者保有分）だと株価収益率の水準が
        # 論文の乗数表から逆算した値より約2割低くなる（PERR 誤差 0.42→0.18、資本コスト 0.59→0.36）
        i_liab = next(i for i, r in enumerate(rows) if r and r[0] is not None and "３．負債" in str(r[0]).replace("　", ""))
        rec[y]["SHARETV"] = float(_find(rows[i_liab:], "うち株式")[8])
        rec[y]["SHARETV_ASSET"] = float(_find(rows, "うち株式")[8])
    # 対外資産・負債残高（暦年末）
    rows = list(openpyxl.load_workbook(VT.annual_file("ss5"), read_only=True, data_only=True)
                ["対外資産・負債残高"].iter_rows(values_only=True))
    years = rows[5]
    for key, lab in {"FASSTV": "対外資産", "FLIABV": "対外負債", "SBCV": "対外純資産"}.items():
        r = next(r for r in rows if r[0] is not None and str(r[0]).strip() == lab)
        for c in range(1, len(r)):
            if years[c] is not None and r[c] not in (None, "-"):
                rec.setdefault(int(years[c]), {})[key] = float(r[c])
    # 家計（個人企業を含む）期末貸借対照表: 家計保有の土地・株式（暦年末）
    rows = list(openpyxl.load_workbook(VT.annual_file("si4"), read_only=True, data_only=True)
                ["期末貸借対照表"].iter_rows(values_only=True))
    years = rows[7]
    land, share = _find(rows, "ａ．土地"), _find(rows, "うち株式")  # 株式は資産側（最初の出現）
    for c in range(1, len(years)):
        if years[c] is not None and str(years[c]).strip().isdigit():
            y = int(str(years[c]).strip())
            rec.setdefault(y, {})["LANDV_HH"] = float(land[c])
            rec[y]["SHAREV_HH"] = float(share[c])
    # 一般政府の固定資産の純取得（年度、GFS）
    rows = list(openpyxl.load_workbook(VT.annual_file("s6_2"), read_only=True, data_only=True)
                ["経常・資本取引"].iter_rows(values_only=True))
    # 年度見出しは各5列ブロック（中央・地方・社保・部門間調整・一般政府）の2列目に置かれている
    # GFS の「純取得」は固定資本減耗控除後なので、総固定資本形成 = 311 固定資産 + 23 固定資本減耗
    gf, cca = _find(rows, "311 固定資産"), _find(rows, "23 固定資本減耗")
    # 一般政府の純貸出(+)／純借入(-)（年度）。論文の財政バランス BGV（出所 CAO,SNA）に当たる
    nl = _find(rows, "純貸出")
    for c, v in enumerate(rows[2]):
        if v and "（" in str(v):
            fy = int(str(v).split("（")[1][:4])
            rec.setdefault(fy, {})["GOVGFCF_FY"] = float(gf[c + 3]) + float(cca[c + 3])
            rec[fy]["GOVNL_FY"] = float(nl[c + 3])
    return pd.DataFrame.from_dict(rec, orient="index").sort_index()


def main() -> None:
    print(f"SNA の版: {VT.NAME} — {VT.V['label']}")
    df = build()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(OUT, encoding="utf-8-sig", index_label="period")
    ann = annual_block()
    ann.to_csv(VT.processed("sna_annual.csv"), encoding="utf-8-sig", index_label="year")
    print(ann.loc[2014:2022].round(0).to_string())
    show = ["GDP", "GDPV", "PGDP", "CP", "IFP", "IG", "YWV", "YCV", "YDV", "TYPV", "TYCV",
            "TCIV", "CSSV", "BSSV", "CCAV", "OTYDV", "RTRIV", "PTRIV"]
    print(f"{OUT} に {df.shape} を書き出しました")
    print(df.loc["2017Q1":"2020Q4", show].round(1).to_string())


if __name__ == "__main__":
    main()
